"""
Wandelt die Hilfeartikel Wissen/*.md in Word-Dateien (Wissen/docx/*.docx) für die
SharePoint-Bibliothek "IT-Hilfe" um. Unterstützt bewusst nur das, was die Artikel nutzen:
# / ## Überschriften, nummerierte und Aufzählungslisten, **fett**, Absätze.
Platzhalter {…} werden gelb markiert, damit sie vor dem Hochladen auffallen.

Aufruf:  python md_zu_docx.py
"""
import re
from pathlib import Path

import docx.shared
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ORDNER = Path(__file__).parent / "Wissen"
ZIEL = ORDNER / "docx"


def text_mit_format(absatz, text: str) -> None:
    # **fett** und {Platzhalter} erkennen, Rest normal
    for teil in re.split(r"(\*\*[^*]+\*\*|\{[^}]+\})", text):
        if not teil:
            continue
        if teil.startswith("**") and teil.endswith("**"):
            absatz.add_run(teil[2:-2]).bold = True
        elif teil.startswith("{") and teil.endswith("}"):
            lauf = absatz.add_run(teil)
            lauf.font.highlight_color = WD_COLOR_INDEX.YELLOW
        else:
            absatz.add_run(teil)


def umwandeln(md: Path) -> Path:
    dok = Document()
    for zeile in md.read_text(encoding="utf-8").splitlines():
        z = zeile.rstrip()
        if not z:
            continue
        if z.startswith("# "):
            dok.add_heading(z[2:], level=1)
        elif z.startswith("## "):
            dok.add_heading(z[3:], level=2)
        elif re.match(r"^\d+\.\s", z):
            # Nummer als Text behalten: Word-"List Number" zählt sonst über alle Listen
            # des Dokuments weiter (Schritt 7, 8 … statt 1, 2 im nächsten Abschnitt)
            absatz = dok.add_paragraph()
            absatz.paragraph_format.left_indent = docx.shared.Cm(0.6)
            text_mit_format(absatz, z)
        elif z.startswith("- "):
            text_mit_format(dok.add_paragraph(style="List Bullet"), z[2:])
        else:
            text_mit_format(dok.add_paragraph(), z)
    ZIEL.mkdir(exist_ok=True)
    ziel = ZIEL / (md.stem + ".docx")
    dok.save(ziel)
    return ziel


if __name__ == "__main__":
    for datei in sorted(ORDNER.glob("*.md")):
        print("OK", umwandeln(datei).name)
